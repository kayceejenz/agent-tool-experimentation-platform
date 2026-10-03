"""Read bundled Excel tables and atomically append demo records to a working file.

Source files are immutable. The working workbook is the only mutation boundary;
keep it closed in Excel while the server writes. This is a local demo, not an
Excel replacement for a production transactional database.
"""

import hashlib
import json
import os
import tempfile
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from filelock import FileLock
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

SEED = Path(__file__).parent / "seed"
SOURCES = {
    "products": ("products.xlsx", "Products"),
    "users": ("users.xlsx", "Users"),
    "orders": ("orders.xlsx", "Orders"),
    "order_items": ("orders.xlsx", "OrderItems"),
    "purchase_history": ("purchase_history.xlsx", "Purchases"),
    "recommendations": ("recommendations.xlsx", "Recommendations"),
}


def value(cell):
    if isinstance(cell, (datetime, date)):
        return cell.isoformat()[:10]
    return cell


def records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [
        dict(zip(headers, [value(cell) for cell in row], strict=True))
        for row in rows
        if any(cell is not None for cell in row)
    ]


class ExcelStore:
    def __init__(self, directory: Path, seed: Path = SEED):
        self.seed = seed
        self.directory = directory.resolve()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.path = self.directory / "records.xlsx"
        self.lock = FileLock(str(self.directory / "records.lock"), timeout=10)

    def source(self, table):
        filename, sheet = SOURCES[table]
        workbook = load_workbook(self.seed / filename, read_only=True, data_only=True)
        try:
            return records(workbook[sheet])
        finally:
            workbook.close()

    def working(self, table):
        with self.lock:
            path = self.path if self.path.exists() else self.seed / "records.xlsx"
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return records(workbook[table])
            finally:
                workbook.close()

    def rows(self, table):
        rows = self.source(table)
        if table in ("users", "orders", "order_items"):
            rows += self.working(
                {"users": "Users", "orders": "Orders", "order_items": "OrderItems"}[
                    table
                ]
            )
        return rows

    @contextmanager
    def transaction(self):
        with self.lock:
            workbook = load_workbook(
                self.path if self.path.exists() else self.seed / "records.xlsx"
            )
            try:
                yield workbook
                with tempfile.NamedTemporaryFile(
                    dir=self.directory, suffix=".xlsx", delete=False
                ) as output:
                    temporary = Path(output.name)
                try:
                    workbook.save(temporary)
                    with temporary.open("rb") as handle:
                        os.fsync(handle.fileno())
                    os.replace(temporary, self.path)
                finally:
                    temporary.unlink(missing_ok=True)
            finally:
                workbook.close()

    @staticmethod
    def append(workbook, sheet_name, row):
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        while sheet.max_row > 1 and all(
            cell.value is None for cell in sheet[sheet.max_row]
        ):
            sheet.delete_rows(sheet.max_row)
        sheet.append([row.get(header) for header in headers])
        for cell in sheet[sheet.max_row]:
            # Force caller-supplied text to literal strings, never Excel formulas.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, color="172033")
            if sheet.max_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EFF6FF")
            field = headers[cell.column - 1]
            if field.endswith("_date") and cell.value:
                cell.value = datetime.fromisoformat(str(cell.value))
                cell.number_format = "dd mmm yyyy"
            elif field in ("price", "unit_price", "line_total", "total"):
                cell.number_format = '"£"#,##0.00'
            elif field == "discount":
                cell.number_format = "0.0%"
            if cell.column_letter in ["A", "B"]:
                cell.number_format = "@"
        for table in sheet.tables.values():
            table.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}{sheet.max_row}"

    @staticmethod
    def request(workbook, key, operation, payload):
        digest = hashlib.sha256(
            json.dumps(payload, sort_keys=True).encode()
        ).hexdigest()
        for row in records(workbook["Requests"]):
            if row["request_key"] == key:
                if row["operation"] != operation or row["payload_hash"] != digest:
                    raise ValueError(
                        "This request key was already used with different details."
                    )
                return row["entity_id"], digest
        return None, digest
